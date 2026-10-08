"""
generador.py — Motor de generación de plantillas XLSM.
Unifica la lógica de NORTE y CENTRO en una sola función parametrizable
según ZONA_CONFIG definido en config.py.
"""
import re
import os
import shutil
import unicodedata
from decimal import Decimal
from datetime import datetime
from copy import copy
from typing import Callable, Optional, List, Dict, Tuple
from openpyxl import load_workbook

from config import ZONA_CONFIG, PLANTILLA_START_ROW
from modules.utils import (
    normalize, excel_round,
    match_keyword, match_keyword_loose,
    SPECIAL_ITEMS,
)

# ─── Descripciones PM03 y mapeo de línea ──────────────────────────────────────
PM03_DESCS = {
    "E":  "Srv.mtto Preventivo de Equipo de Bombeo en CR",
    "TA": "Srv.mtto Preventivo de Tableros Electricos en CR",
    "AH": "Srv.mtto Preventivo de Accesorios Hidraulicos en CR",
}
LINEA_MAP = {"E": 1, "TA": 2, "AH": 3}


# ─── ACT code según número de equipos ─────────────────────────────────────────
def get_act_code(num_equipos: int) -> str:
    if num_equipos <= 2:   return "10003486"
    elif num_equipos <= 4: return "10003487"
    elif num_equipos <= 6: return "10003488"
    elif num_equipos <= 8: return "10003489"
    else:                  return "10003490"


# ─── Parseo del nombre de archivo ─────────────────────────────────────────────
def parse_filename(fname: str):
    m_cr = re.search(r"CR-(\d+[A-Za-z]?)", fname)
    if not m_cr:
        return None, None
    cr_base = m_cr.group(1)
    variant = ""
    if "500" in fname.upper() and "HP" in fname.upper():
        variant = "_500HP"
    elif "200" in fname.upper() and "HP" in fname.upper():
        variant = "_200HP"
    cr_code = cr_base + variant

    m_date = re.search(r"(\d{2})[._](\d{2})[._](\d{4})", fname)
    if not m_date:
        return cr_code, None
    try:
        fecha = datetime(int(m_date.group(3)), int(m_date.group(2)), int(m_date.group(1)))
    except Exception:
        fecha = None
    return cr_code, fecha


# ─── Catálogo (MATERIALES + ACTIVIDAD) ────────────────────────────────────────
def load_catalog(wb_estilos) -> dict:
    catalog: dict = {}

    ws_mat = wb_estilos["MATERIALES"]
    for row in ws_mat.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            break
        sap  = str(row[0]).strip()
        desc = str(row[2]).strip() if row[2] else ""
        try:
            precio = excel_round(float(row[4])) if row[4] else Decimal("0")
        except Exception:
            precio = Decimal("0")
        catalog[sap] = {"desc": desc, "precio": precio}

    ws_act = wb_estilos["ACTIVIDAD"]
    for row in ws_act.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            break
        sap  = str(row[0]).strip()
        desc = str(row[2]).strip() if row[2] else ""
        try:
            precio = excel_round(float(row[3])) if row[3] else Decimal("0")
        except Exception:
            precio = Decimal("0")
        catalog[sap] = {"desc": desc, "precio": precio}

    for token, info in SPECIAL_ITEMS.items():
        catalog[token] = {"desc": info["desc"], "precio": info["precio"]}

    return catalog


# ─── Estilos ──────────────────────────────────────────────────────────────────
def copy_cell_style(src_cell) -> dict:
    style = {}
    if src_cell.has_style:
        style["font"]          = copy(src_cell.font)
        style["fill"]          = copy(src_cell.fill)
        style["border"]        = copy(src_cell.border)
        style["alignment"]     = copy(src_cell.alignment)
        style["number_format"] = src_cell.number_format
    return style


def apply_style(dst_cell, style: dict) -> None:
    if "font"          in style: dst_cell.font          = copy(style["font"])
    if "fill"          in style: dst_cell.fill          = copy(style["fill"])
    if "border"        in style: dst_cell.border        = copy(style["border"])
    if "alignment"     in style: dst_cell.alignment     = copy(style["alignment"])
    if "number_format" in style: dst_cell.number_format = style["number_format"]


def load_styles(wb_estilos, cfg: dict) -> dict:
    ws = wb_estilos["PLANTILLA MODELO"]
    num_cols = cfg["num_cols"]
    ref_rows = cfg["ref_rows"]

    styles: dict = {}
    for nombre, fila in ref_rows.items():
        row_styles = []
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=fila, column=col)
            row_styles.append(copy_cell_style(cell))
        styles[nombre] = row_styles

    styles["TA_MAT"] = styles["MAT"]
    styles["AH_MAT"] = styles["MAT"]
    return styles


# ─── Parseo del archivo fuente ────────────────────────────────────────────────
def parse_source_file(filepath: str, cr_code: str, cfg: dict) -> List[dict]:
    """
    Lee el archivo fuente y devuelve lista de equipos con sus materiales E/TA/AH.
    """
    header_row  = cfg["header_row"]
    tipo_row    = cfg["tipo_row"]
    data_start  = cfg["data_start"]

    wb = load_workbook(filepath, data_only=True)
    ws = wb.active

    cr_base = re.match(r"(\d+[A-Za-z]?)", cr_code).group(1)
    cr_pattern = re.compile(r"CR[\s\-]*" + re.escape(cr_base) + r"\b", re.IGNORECASE)

    # Leer filas de cabecera
    h_headers: dict[int, str] = {}
    for cell in ws[header_row]:
        if cell.value:
            h_headers[cell.column] = str(cell.value)

    h_tipos: dict[int, str] = {}
    for cell in ws[tipo_row]:
        if cell.value:
            h_tipos[cell.column] = str(cell.value)

    # Detectar columnas relevantes (BOMBA + adyacentes TABLERO/ACCES)
    relevant_cols: dict[int, dict] = {}
    for col, text in h_headers.items():
        if not cr_pattern.search(text):
            continue
        m_eq = re.search(r"n[°o]?\s*(\d+)", text, re.IGNORECASE)
        if not m_eq:
            continue
        eq_num = int(m_eq.group(1))

        for dc in range(0, 4):
            next_col = col + dc
            if dc > 0 and next_col in h_headers:
                break
            tipo_txt  = h_tipos.get(next_col, "")
            tipo_norm = normalize(tipo_txt)
            if not tipo_txt:
                if dc == 0:
                    relevant_cols[next_col] = {"equipo": eq_num, "tipo": "E"}
                break
            if "bomb" in tipo_norm or "equip" in tipo_norm:
                tipo = "E"
            elif "tablero" in tipo_norm or "electr" in tipo_norm:
                tipo = "TA"
            elif "acces" in tipo_norm or "hidra" in tipo_norm:
                tipo = "AH"
            else:
                break
            relevant_cols[next_col] = {"equipo": eq_num, "tipo": tipo}

    if not relevant_cols:
        print(f"  ADVERTENCIA: Sin columnas para CR-{cr_base} en {os.path.basename(filepath)}")
        wb.close()
        return []

    # Leer materiales desde data_start
    equipos: dict[int, dict] = {}
    SKIP = {"DESCRIPCION", "DESCRIPCIÓN", "MATERIAL", "TOTAL"}

    for row in ws.iter_rows(min_row=data_start, values_only=True):
        desc_val = row[1] if len(row) > 1 else None
        if not desc_val:
            continue
        desc_str  = str(desc_val).strip()
        if not desc_str or desc_str.upper() in SKIP:
            continue
        desc_norm = normalize(desc_str)
        if any(x in desc_norm for x in (
            "equipo de bombeo", "tablero elect", "accesorios hidr",
            "bomba n", "tablero n", "accesorio n",
        )):
            continue

        for col, info in relevant_cols.items():
            idx = col - 1
            if idx >= len(row):
                continue
            qty_val = row[idx]
            if qty_val is None or qty_val == "" or qty_val == 0:
                continue
            try:
                qty = excel_round(float(qty_val))
            except (ValueError, TypeError):
                continue
            if qty <= 0:
                continue

            eq_num = info["equipo"]
            tipo   = info["tipo"]
            if eq_num not in equipos:
                equipos[eq_num] = {"E": [], "TA": [], "AH": []}
            equipos[eq_num][tipo].append({"desc": desc_str, "qty": qty})

    wb.close()
    result = []
    for eq_num in sorted(equipos.keys()):
        d = equipos[eq_num]
        if d["E"] or d["TA"] or d["AH"]:
            result.append({"equipo": eq_num, "E": d["E"], "TA": d["TA"], "AH": d["AH"]})
    return result


# ─── Resolver material a SAP / precio ─────────────────────────────────────────
def resolver_material(mat: dict, catalog: dict, unmapped_log: list, tag: str) -> dict:
    sap_token = match_keyword(mat["desc"]) or match_keyword_loose(mat["desc"])
    if sap_token is None:
        unmapped_log.append(f"{tag}: '{mat['desc']}'")
        sap_code, desc_final, precio = None, mat["desc"], Decimal("0")
    elif sap_token in SPECIAL_ITEMS:
        sp = SPECIAL_ITEMS[sap_token]
        sap_code = sp.get("sap")
        desc_final = sp["desc"]
        precio     = sp["precio"]
    else:
        info = catalog.get(sap_token, {"desc": mat["desc"], "precio": Decimal("0")})
        sap_code, desc_final, precio = sap_token, info["desc"], info["precio"]
    subtotal = excel_round(precio * mat["qty"])
    return {
        "sap": sap_code, "desc": desc_final,
        "qty": mat["qty"], "precio": precio, "subtotal": subtotal,
    }


# ─── Construcción de filas ────────────────────────────────────────────────────
def build_cr_rows(
    cr_base: str, fecha_str: str,
    equipos_data: List[dict],
    catalog: dict,
    unmapped_log: list,
) -> List[dict]:
    rows = []
    num_equipos = len(equipos_data)
    act_sap     = get_act_code(num_equipos)
    act_info    = catalog.get(act_sap, {"desc": "MP en CR", "precio": Decimal("0")})
    act_precio  = act_info["precio"]
    cr_display  = f"CR-{cr_base}"
    first_equipo = True

    for eq_data in equipos_data:
        eq_num = eq_data["equipo"]

        # ── BOMBA (E) ──────────────────────────────────────────────────────────
        e_mats = [
            resolver_material(m, catalog, unmapped_log, f"CR-{cr_base} E{eq_num}")
            for m in eq_data["E"]
        ]
        e_sub_mats = sum((m["subtotal"] for m in e_mats), Decimal("0"))
        e_pm03_sub = (e_sub_mats + act_precio) if first_equipo else e_sub_mats

        rows.append(_row(f"E{eq_num}", cr_display, "E", "PM03", None,
                         PM03_DESCS["E"], 1, e_pm03_sub, e_pm03_sub, fecha_str, "E_PM03"))
        if first_equipo:
            rows.append(_row(f"E{eq_num}", cr_display, "E", "ACT", act_sap,
                             act_info["desc"], 1, act_precio, act_precio, fecha_str, "ACT"))
        for mat in e_mats:
            rows.append(_row(f"E{eq_num}", cr_display, "E", "MAT", mat["sap"],
                             mat["desc"], mat["qty"], mat["precio"], mat["subtotal"],
                             fecha_str, "MAT"))

        # ── TABLERO (TA) ───────────────────────────────────────────────────────
        ta_mats = [
            resolver_material(m, catalog, unmapped_log, f"CR-{cr_base} TA{eq_num}")
            for m in eq_data["TA"]
        ]
        ta_sub = sum((m["subtotal"] for m in ta_mats), Decimal("0"))
        rows.append(_row(f"TA{eq_num}", cr_display, "TA", "PM03", None,
                         PM03_DESCS["TA"], 1, ta_sub, ta_sub, fecha_str, "TA_PM03"))
        for mat in ta_mats:
            rows.append(_row(f"TA{eq_num}", cr_display, "TA", "MAT", mat["sap"],
                             mat["desc"], mat["qty"], mat["precio"], mat["subtotal"],
                             fecha_str, "TA_MAT"))

        # ── ACCESORIOS (AH) ────────────────────────────────────────────────────
        ah_mats = [
            resolver_material(m, catalog, unmapped_log, f"CR-{cr_base} AH{eq_num}")
            for m in eq_data["AH"]
        ]
        ah_sub = sum((m["subtotal"] for m in ah_mats), Decimal("0"))
        rows.append(_row(f"AH{eq_num}", cr_display, "AH", "PM03", None,
                         PM03_DESCS["AH"], 1, ah_sub, ah_sub, fecha_str, "AH_PM03"))
        for mat in ah_mats:
            rows.append(_row(f"AH{eq_num}", cr_display, "AH", "MAT", mat["sap"],
                             mat["desc"], mat["qty"], mat["precio"], mat["subtotal"],
                             fecha_str, "AH_MAT"))

        first_equipo = False

    return rows


def _row(orden, ubicacion, linea_key, tipo, sap, desc, qty, precio, subtotal, fecha, row_type):
    return {
        "orden": orden, "operacion": 10,
        "ubicacion": ubicacion, "linea": LINEA_MAP[linea_key],
        "tipo": tipo, "sap": sap, "desc": desc,
        "qty": Decimal(str(qty)), "precio": precio, "subtotal": subtotal,
        "fecha": fecha, "row_type": row_type,
    }


# ─── Escritura en plantilla ───────────────────────────────────────────────────
def write_row_to_sheet(ws, row_num: int, row_data: dict, styles: dict, num_cols: int):
    row_type  = row_data["row_type"]
    style_map = {
        "E_PM03": "E_PM03", "ACT": "ACT",
        "TA_PM03": "TA_PM03", "AH_PM03": "AH_PM03",
        "MAT": "MAT", "TA_MAT": "TA_MAT", "AH_MAT": "AH_MAT",
    }
    style_key  = style_map.get(row_type, "MAT")
    row_styles = styles.get(style_key, [{}] * num_cols)

    # Primeras 11 columnas de datos
    data_cols = [
        row_data["orden"],           # C1
        10,                          # C2  operación
        row_data["ubicacion"],       # C3
        row_data["linea"],           # C4
        row_data["tipo"],            # C5
        row_data["sap"],             # C6
        row_data["desc"],            # C7
        float(row_data["qty"]),      # C8
        float(row_data["precio"]),   # C9
        float(row_data["subtotal"]), # C10
        row_data["fecha"],           # C11
    ]
    # Padding para zonas con más columnas (NORTE: 21)
    values = data_cols + [None] * (num_cols - len(data_cols))

    for col_idx, (val, col_style) in enumerate(zip(values, row_styles), start=1):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.value = val
        if col_style:
            apply_style(cell, col_style)


# ─── Función principal ────────────────────────────────────────────────────────
def generar_plantilla(
    zona: str,
    source_paths: list[str],
    plantilla_path: str,
    estilos_path: str,
    output_path: str,
    log_callback: Optional[Callable] = None,
) -> tuple[list[str], int]:
    """
    Genera el archivo XLSM de salida a partir de los archivos fuente.

    Args:
        zona:           "NORTE" o "CENTRO"
        source_paths:   Lista de rutas a archivos .xlsx fuente
        plantilla_path: Ruta a la plantilla vacía .xlsm
        estilos_path:   Ruta al archivo de estilos .xlsm
        output_path:    Ruta destino del archivo generado
        log_callback:   Función opcional (mensaje: str, progreso: float) → None

    Returns:
        (logs: list[str], total_rows: int)
    """
    cfg  = ZONA_CONFIG[zona.upper()]
    logs: list[str] = []
    n    = len(source_paths)

    def log(msg: str, pct: float = 0.0):
        logs.append(msg)
        if log_callback:
            log_callback(msg, pct)
        print(msg)

    log(f"▶ Iniciando generación {zona}  ({n} archivos fuente)", 0.0)

    # 1. Catálogo y estilos
    log("[1/4] Cargando catálogo y estilos...", 0.05)
    wb_estilos = load_workbook(estilos_path, keep_vba=True, data_only=True)
    catalog = load_catalog(wb_estilos)
    styles  = load_styles(wb_estilos, cfg)
    wb_estilos.close()
    log(f"  Catálogo: {len(catalog)} entradas", 0.10)

    # 2. Parseo de archivos fuente
    log("[2/4] Parseando archivos fuente...", 0.10)
    cr_data_list: List[dict] = []
    unmapped_log: list[str]  = []

    for i, filepath in enumerate(sorted(source_paths)):
        pct = 0.10 + 0.60 * i / max(n, 1)
        fname    = os.path.basename(filepath)
        cr_code, fecha = parse_filename(fname)

        if cr_code is None:
            log(f"  IGNORADO (sin código CR): {fname}", pct)
            continue

        fecha_str = fecha.strftime("%d/%m/%Y") if fecha else ""
        log(f"  [{i+1}/{n}] CR-{cr_code} ({fecha_str})", pct)

        equipos = parse_source_file(filepath, cr_code, cfg)
        for eq in equipos:
            log(
                f"    N°{eq['equipo']}: "
                f"E={len(eq['E'])} TA={len(eq['TA'])} AH={len(eq['AH'])} mats",
                pct,
            )

        if not equipos:
            log(f"    ⚠ Sin datos para CR-{cr_code}", pct)
            continue

        base = re.match(r"(\d+[A-Za-z]?)", cr_code).group(1)
        rows = build_cr_rows(base, fecha_str, equipos, catalog, unmapped_log)
        cr_data_list.append({"cr": cr_code, "rows": rows})
        log(f"    → {len(rows)} filas generadas", pct)

    total_filas = sum(len(c["rows"]) for c in cr_data_list)
    log(f"[3/4] Escribiendo {total_filas} filas en plantilla...", 0.72)

    # 3. Escribir en plantilla
    shutil.copy2(plantilla_path, output_path)
    wb_out  = load_workbook(output_path, keep_vba=True)
    ws_out  = wb_out["PLANTILLA MODELO"]
    cur_row = PLANTILLA_START_ROW

    for cr_entry in cr_data_list:
        for row_data in cr_entry["rows"]:
            write_row_to_sheet(ws_out, cur_row, row_data, styles, cfg["num_cols"])
            cur_row += 1

    # 4. Guardar
    log("[4/4] Guardando...", 0.95)
    wb_out.save(output_path)
    wb_out.close()

    # Resumen de no mapeados
    if unmapped_log:
        vistos: set[str] = set()
        log(f"\n⚠ MATERIALES SIN MAPEAR ({len(set(unmapped_log))}):", 0.98)
        for entry in unmapped_log:
            if entry not in vistos:
                log(f"  - {entry}", 0.98)
                vistos.add(entry)
    else:
        log("✓ Todos los materiales mapeados correctamente.", 0.98)

    log(
        f"\n✅ COMPLETADO — {total_filas} filas en {len(cr_data_list)} CRs → {output_path}",
        1.0,
    )
    return logs, total_filas
