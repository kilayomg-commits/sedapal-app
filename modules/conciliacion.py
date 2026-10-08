from typing import Optional, List, Dict, Tuple
"""
conciliacion.py — Lee los totales del archivo XLSM generado y los compara
con los montos teóricos/esperados ingresados por el usuario.
"""
from openpyxl import load_workbook


# ─── Lectura del archivo generado ─────────────────────────────────────────────
def leer_totales_output(filepath: str) -> dict:
    """
    Lee 'PLANTILLA MODELO' del archivo .xlsm y suma los subtotales PM03
    agrupados por CR y tipo (E, TA, AH).

    Estructura de columnas esperada en la hoja:
        A=Orden (E1/TA1/AH1)  C=Ubicacion (CR-XX)
        E=Tipo (PM03/ACT/MAT) J=Subtotal

    Returns:
        {
            'CR-78':  {'E': 1234.56, 'TA': 234.50, 'AH': 123.00, 'total': 1592.06},
            'CR-254': {...},
            ...
        }
    """
    wb = load_workbook(filepath, data_only=True)
    ws = wb["PLANTILLA MODELO"]

    totales: Dict[str, dict] = {}

    for row in ws.iter_rows(min_row=10, values_only=True):
        # Col C (índice 2) = ubicación
        if not row[2]:
            continue
        cr = str(row[2]).strip()
        if not cr.startswith("CR-"):
            continue

        tipo_fila = str(row[4]).strip() if row[4] else ""  # Col E
        if tipo_fila != "PM03":
            continue

        orden    = str(row[0]).strip() if row[0] else ""   # Col A
        subtotal = float(row[9]) if row[9] is not None else 0.0  # Col J

        if cr not in totales:
            totales[cr] = {"E": 0.0, "TA": 0.0, "AH": 0.0}

        orden_upper = orden.upper()
        if orden_upper.startswith("TA"):
            totales[cr]["TA"] += subtotal
        elif orden_upper.startswith("AH"):
            totales[cr]["AH"] += subtotal
        elif orden_upper.startswith("E"):
            totales[cr]["E"] += subtotal

    wb.close()

    for cr in totales:
        totales[cr]["total"] = (
            totales[cr]["E"] + totales[cr]["TA"] + totales[cr]["AH"]
        )

    return totales


# ─── Comparación ──────────────────────────────────────────────────────────────
def comparar_con_esperado(
    totales_reales: dict,
    esperados: dict,
) -> List[dict]:
    """
    Genera la tabla de conciliación comparando totales reales vs esperados.

    Args:
        totales_reales: salida de leer_totales_output()
        esperados:      {'CR-78': 5000.0, 'CR-254': 3200.0, ...}

    Returns:
        Lista de dicts listos para convertir en DataFrame de Streamlit.
    """
    resultados = []
    for cr in sorted(totales_reales.keys()):
        real     = totales_reales[cr]["total"]
        esperado = esperados.get(cr, 0.0)
        dif      = real - esperado
        var_pct  = ((dif / esperado) * 100) if esperado != 0 else 0.0

        if abs(dif) < 0.02:
            estado = "✅ OK"
        elif dif > 0:
            estado = "⬆️ Sobre"
        else:
            estado = "⬇️ Bajo"

        resultados.append({
            "CR":            cr,
            "E (S/)":        round(totales_reales[cr]["E"],  2),
            "TA (S/)":       round(totales_reales[cr]["TA"], 2),
            "AH (S/)":       round(totales_reales[cr]["AH"], 2),
            "Real (S/)":     round(real,     2),
            "Esperado (S/)": round(esperado, 2),
            "Diferencia":    round(dif,      2),
            "Var %":         round(var_pct,  1),
            "Estado":        estado,
        })

    return resultados
